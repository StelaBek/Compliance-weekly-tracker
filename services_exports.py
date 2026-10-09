from __future__ import annotations

from io import BytesIO
from textwrap import wrap
import re

import pandas as pd
from reportlab.pdfgen import canvas
from reportlab.lib.colors import HexColor
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from core_analytics import calculate_metrics, executive_summary, deadline_frame, impact_counts

PURPLE = "#A20B8D"
YELLOW = "#FFC400"
LIGHT_BG = "#FFFFFF"
LIGHT_SURFACE = "#F6F7FA"
LIGHT_TEXT = "#111827"
LIGHT_MUTED = "#6B7280"
LIGHT_BORDER = "#E5E7EB"
DARK_BG = "#0E1117"
DARK_SURFACE = "#171B22"
DARK_TEXT = "#F5F7FA"
DARK_MUTED = "#A4ACB8"
DARK_BORDER = "#303743"

IMPACT_ORDER = ["Informational", "Low", "Medium", "High", "Critical"]


def _text(value, fallback="—"):
    if value is None or (isinstance(value, float) and pd.isna(value)) or str(value).strip() == "":
        return fallback
    return str(value)


def _clean_ai(text: str) -> str:
    text = str(text or "").strip()
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"^#+\s*", "", text, flags=re.MULTILINE)
    return text


def _theme(theme: str):
    dark = str(theme or "").lower() == "dark"
    return {
        "bg": DARK_BG if dark else LIGHT_BG,
        "surface": DARK_SURFACE if dark else LIGHT_SURFACE,
        "text": DARK_TEXT if dark else LIGHT_TEXT,
        "muted": DARK_MUTED if dark else LIGHT_MUTED,
        "border": DARK_BORDER if dark else LIGHT_BORDER,
        "dark": dark,
    }


def _short(value, n=68):
    s = _text(value, "")
    return s if len(s) <= n else s[: n - 1] + "…"


def _pdf_text(c, text, x, y, max_chars, font="Helvetica", size=8, color="#111827", leading=None, max_lines=2):
    leading = leading or size * 1.25
    c.setFillColor(HexColor(color))
    c.setFont(font, size)
    lines = []
    for paragraph in str(text or "").splitlines() or [""]:
        lines.extend(wrap(paragraph, max_chars) or [""])
    lines = lines[:max_lines]
    for i, line in enumerate(lines):
        if i == max_lines - 1 and len(lines) == max_lines and len(line) >= max_chars:
            line = line[: max(1, max_chars - 1)] + "…"
        c.drawString(x, y - i * leading, line)
    return y - len(lines) * leading


def _pdf_box(c, x, y, w, h, fill, border, radius=8):
    c.setFillColor(HexColor(fill))
    c.setStrokeColor(HexColor(border))
    c.roundRect(x, y, w, h, radius, fill=1, stroke=1)


def _pdf_dashboard_page(c, df: pd.DataFrame, ai: str, theme: str):
    t = _theme(theme)
    W, H = 960, 540
    c.setFillColor(HexColor(t["bg"]))
    c.rect(0, 0, W, H, fill=1, stroke=0)

    # Header mirrors the visible web hero.
    _pdf_box(c, 26, 466, 908, 54, t["surface"], t["border"], 10)
    c.setFillColor(HexColor(PURPLE))
    c.setFont("Helvetica-Bold", 8)
    c.drawString(42, 503, "EUROPEAN PRODUCT & TAX COMPLIANCE")
    c.setFillColor(HexColor(t["text"]))
    c.setFont("Helvetica-Bold", 22)
    c.drawString(42, 479, "Regulatory intelligence")

    m = calculate_metrics(df)
    metrics = [
        ("Relevant findings", m["total"]),
        ("High / critical", m["high_critical"]),
        ("Deadlines ≤90d", m["next_90_days"]),
        ("Jurisdictions", m["jurisdictions"]),
    ]
    mx, my, mw, mh, gap = 26, 410, 218, 44, 12
    for i, (label, value) in enumerate(metrics):
        x = mx + i * (mw + gap)
        _pdf_box(c, x, my, mw, mh, t["surface"], t["border"], 8)
        c.setFillColor(HexColor(t["muted"]))
        c.setFont("Helvetica", 7)
        c.drawString(x + 12, my + 28, label)
        c.setFillColor(HexColor(PURPLE))
        c.setFont("Helvetica-Bold", 15)
        c.drawString(x + 12, my + 10, str(value))

    # Executive summary.
    c.setFillColor(HexColor(t["text"]))
    c.setFont("Helvetica-Bold", 11)
    c.drawString(26, 386, "Executive summary")
    c.setFillColor(HexColor(t["muted"]))
    c.setFont("Helvetica", 6.5)
    c.drawString(26, 374, "Python-generated · deterministic from the current filtered dataset")
    _pdf_text(c, executive_summary(df), 26, 359, 145, size=7.2, color=t["text"], max_lines=3)

    # Left: developments table.
    left_x, left_y, left_w, left_h = 26, 154, 548, 166
    c.setFillColor(HexColor(t["text"]))
    c.setFont("Helvetica-Bold", 11)
    c.drawString(left_x, 334, "Priority developments")
    _pdf_box(c, left_x, left_y, left_w, left_h, t["surface"], t["border"], 8)
    headers = [("Regulatory scope", 112), ("Development", 280), ("Impact", 74)]
    hx = left_x + 10
    for label, width in headers:
        c.setFillColor(HexColor(t["muted"]))
        c.setFont("Helvetica-Bold", 6.2)
        c.drawString(hx, left_y + left_h - 16, label)
        hx += width
    row_y = left_y + left_h - 32
    if df.empty:
        _pdf_text(c, "No live findings match the current filters yet.", left_x + 10, row_y, 90, size=7.2, color=t["muted"], max_lines=2)
    else:
        for _, r in df.head(7).iterrows():
            c.setStrokeColor(HexColor(t["border"]))
            c.line(left_x + 8, row_y - 4, left_x + left_w - 8, row_y - 4)
            c.setFillColor(HexColor(t["text"]))
            c.setFont("Helvetica", 6.5)
            c.drawString(left_x + 10, row_y + 3, _short(_scope_text(r), 22))
            c.drawString(left_x + 124, row_y + 3, _short(r.get("english_title"), 58))
            c.setFillColor(HexColor(PURPLE))
            c.setFont("Helvetica-Bold", 6.5)
            c.drawString(left_x + 407, row_y + 3, _short(r.get("business_impact"), 12))
            row_y -= 18

    # Right: impact profile.
    right_x, right_w = 598, 336
    c.setFillColor(HexColor(t["text"]))
    c.setFont("Helvetica-Bold", 11)
    c.drawString(right_x, 334, "Impact profile")
    _pdf_box(c, right_x, 226, right_w, 94, t["surface"], t["border"], 8)
    counts = {str(r["Impact"]): int(r["Count"]) for _, r in impact_counts(df).iterrows()} if not df.empty else {}
    max_count = max([counts.get(k, 0) for k in IMPACT_ORDER] + [1])
    chart_x, chart_y, chart_h = right_x + 18, 244, 48
    bar_w, bgap = 35, 25
    for i, impact in enumerate(IMPACT_ORDER):
        val = counts.get(impact, 0)
        x = chart_x + i * (bar_w + bgap)
        h = chart_h * val / max_count
        c.setFillColor(HexColor("#7FC4F2" if impact not in {"High", "Critical"} else YELLOW))
        c.rect(x, chart_y, bar_w, max(1, h), fill=1, stroke=0)
        c.setFillColor(HexColor(t["muted"]))
        c.setFont("Helvetica", 5.2)
        c.drawCentredString(x + bar_w / 2, chart_y - 9, impact)
        if val:
            c.setFillColor(HexColor(t["text"]))
            c.setFont("Helvetica-Bold", 6)
            c.drawCentredString(x + bar_w / 2, chart_y + h + 3, str(val))

    # Right: deadlines.
    c.setFillColor(HexColor(t["text"]))
    c.setFont("Helvetica-Bold", 11)
    c.drawString(right_x, 207, "Next confirmed deadlines")
    _pdf_box(c, right_x, 154, right_w, 42, t["surface"], t["border"], 8)
    deadlines = deadline_frame(df, horizon_days=365)
    if deadlines.empty:
        _pdf_text(c, "No confirmed deadline in the next 12 months for this view.", right_x + 12, 178, 65, size=6.4, color=t["muted"], max_lines=2)
    else:
        first = deadlines.iloc[0]
        _pdf_text(c, f"{_text(first.get('compliance_deadline'))} · {_short(_scope_text(first), 24)} · {_short(first.get('english_title'), 42)}", right_x + 12, 179, 63, size=6.2, color=t["text"], max_lines=2)

    # AI section at the bottom, like the web view.
    c.setStrokeColor(HexColor(t["border"]))
    c.line(26, 132, 934, 132)
    c.setFillColor(HexColor(t["text"]))
    c.setFont("Helvetica-Bold", 11)
    c.drawString(26, 112, "AI analysis")
    c.setFillColor(HexColor(t["muted"]))
    c.setFont("Helvetica", 6.4)
    c.drawString(26, 99, "Optional · grounded in visible structured records · informational only, not legal or tax advice")
    ai_clean = _clean_ai(ai)
    if ai_clean:
        _pdf_box(c, 26, 26, 908, 58, t["surface"], t["border"], 8)
        c.setFillColor(HexColor(PURPLE))
        c.setFont("Helvetica-Bold", 7)
        c.drawString(38, 69, "AI-GENERATED PORTFOLIO ANALYSIS")
        _pdf_text(c, ai_clean, 38, 55, 150, size=6.3, color=t["text"], max_lines=4)
    else:
        c.setFillColor(HexColor(t["muted"]))
        c.setFont("Helvetica", 6.6)
        c.drawString(26, 79, "No AI analysis generated for this view yet.")



def _scope_text(row) -> str:
    level = str(row.get("geographic_level") or "Unknown")
    jurisdiction = str(row.get("jurisdiction") or "Unknown")
    if level == "National":
        return f"National — {jurisdiction}"
    if level == "EU":
        return "EU — European Union"
    return f"{level} — {jurisdiction}"


def _pdf_finding_page(c, row, theme: str):
    t = _theme(theme)
    W, H = 960, 540
    c.setFillColor(HexColor(t["bg"])); c.rect(0, 0, W, H, fill=1, stroke=0)
    _pdf_box(c, 26, 466, 908, 54, t["surface"], t["border"], 10)
    c.setFillColor(HexColor(PURPLE)); c.setFont("Helvetica-Bold", 8); c.drawString(42, 503, "VERIFIED CURRENT-WEEK FINDING")
    c.setFillColor(HexColor(t["text"])); c.setFont("Helvetica-Bold", 16)
    _pdf_text(c, _text(row.get("english_title")), 42, 484, 92, font="Helvetica-Bold", size=13, color=t["text"], max_lines=2)
    fields = [
        ("Jurisdiction", row.get("jurisdiction")), ("Category", row.get("category")),
        ("Publication / update date", row.get("publication_update_date")), ("Effective / application date", row.get("effective_application_date")),
        ("Legislation", row.get("legislation")), ("Status", row.get("status")),
        ("Source", row.get("source_name")), ("Source URL", row.get("source_url")),
    ]
    x1, x2 = 38, 500
    y = 438
    for i, (label, value) in enumerate(fields):
        x = x1 if i % 2 == 0 else x2
        yy = y - (i // 2) * 34
        c.setFillColor(HexColor(t["muted"])); c.setFont("Helvetica", 6.3); c.drawString(x, yy, label.upper())
        _pdf_text(c, _text(value), x, yy-12, 68 if x == x1 else 62, size=7.2, color=t["text"], max_lines=2)
    c.setFillColor(HexColor(t["text"])); c.setFont("Helvetica-Bold", 10); c.drawString(38, 288, "Scope — what products or services does it apply to?")
    _pdf_box(c, 38, 198, 884, 76, t["surface"], t["border"], 8)
    _pdf_text(c, _text(row.get("scope")), 50, 258, 145, size=7, color=t["text"], max_lines=7)
    c.setFont("Helvetica-Bold", 10); c.setFillColor(HexColor(t["text"])); c.drawString(38, 174, "Summary — what changed and why does it matter?")
    _pdf_box(c, 38, 84, 884, 76, t["surface"], t["border"], 8)
    _pdf_text(c, _text(row.get("summary")), 50, 144, 145, size=7, color=t["text"], max_lines=7)
    c.setFont("Helvetica-Bold", 9); c.setFillColor(HexColor(PURPLE)); c.drawString(38, 60, "BUSINESS ACTION")
    _pdf_text(c, _text(row.get("business_action")), 138, 60, 120, size=7, color=t["text"], max_lines=3)


def build_pdf(df: pd.DataFrame, ai_portfolio_analysis: str = "", theme: str = "light") -> bytes:
    """Export the current filtered weekly report, followed by one full page per finding."""
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=(960, 540))
    _pdf_dashboard_page(c, df, ai_portfolio_analysis, theme)
    c.showPage()
    for _, row in df.iterrows():
        _pdf_finding_page(c, row, theme)
        c.showPage()
    c.save()
    return buf.getvalue()


def _rgb(hex_color: str):
    return RGBColor.from_string(hex_color.lstrip("#"))


def _ppt_text(slide, x, y, w, h, text, size=10, bold=False, color=LIGHT_TEXT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = str(text or "")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _ppt_box(slide, x, y, w, h, fill, border, radius=True):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    s.line.color.rgb = _rgb(border)
    s.line.width = Pt(0.7)
    return s


def _ppt_dashboard_slide(slide, df: pd.DataFrame, ai: str, theme: str):
    t = _theme(theme)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = _rgb(t["bg"])

    _ppt_box(slide, .35, .25, 12.63, .75, t["surface"], t["border"])
    _ppt_text(slide, .55, .39, 6, .18, "EUROPEAN PRODUCT & TAX COMPLIANCE", 7, True, PURPLE)
    _ppt_text(slide, .55, .59, 7.5, .32, "Regulatory intelligence", 20, True, t["text"])

    m = calculate_metrics(df)
    metrics = [
        ("Relevant findings", m["total"]),
        ("High / critical", m["high_critical"]),
        ("Deadlines ≤90d", m["next_90_days"]),
        ("Jurisdictions", m["jurisdictions"]),
    ]
    for i, (label, value) in enumerate(metrics):
        x = .35 + i * 3.18
        _ppt_box(slide, x, 1.13, 2.93, .62, t["surface"], t["border"])
        _ppt_text(slide, x + .13, 1.23, 2.3, .18, label, 6.5, False, t["muted"])
        _ppt_text(slide, x + .13, 1.42, 1, .22, value, 14, True, PURPLE)

    _ppt_text(slide, .35, 1.9, 3, .22, "Executive summary", 11, True, t["text"])
    _ppt_text(slide, .35, 2.15, 12, .45, executive_summary(df), 7.2, False, t["text"])

    _ppt_text(slide, .35, 2.72, 4, .22, "Priority developments", 11, True, t["text"])
    _ppt_box(slide, .35, 2.98, 7.42, 2.05, t["surface"], t["border"])
    _ppt_text(slide, .5, 3.1, 1.55, .15, "Regulatory scope", 6, True, t["muted"])
    _ppt_text(slide, 2.05, 3.1, 4.2, .15, "Development", 6, True, t["muted"])
    _ppt_text(slide, 6.35, 3.1, 1.05, .15, "Impact", 6, True, t["muted"])
    if df.empty:
        _ppt_text(slide, .5, 3.42, 6.7, .35, "No live findings match the current filters yet.", 7, False, t["muted"])
    else:
        y = 3.34
        for _, r in df.head(7).iterrows():
            _ppt_text(slide, .5, y, 1.5, .18, _short(_scope_text(r), 22), 6.3, False, t["text"])
            _ppt_text(slide, 2.05, y, 4.15, .18, _short(r.get("english_title"), 58), 6.3, False, t["text"])
            _ppt_text(slide, 6.35, y, 1.05, .18, _short(r.get("business_impact"), 12), 6.3, True, PURPLE)
            y += .23

    _ppt_text(slide, 8.05, 2.72, 3, .22, "Impact profile", 11, True, t["text"])
    _ppt_box(slide, 8.05, 2.98, 4.93, 1.18, t["surface"], t["border"])
    counts = {str(r["Impact"]): int(r["Count"]) for _, r in impact_counts(df).iterrows()} if not df.empty else {}
    max_count = max([counts.get(k, 0) for k in IMPACT_ORDER] + [1])
    for i, impact in enumerate(IMPACT_ORDER):
        val = counts.get(impact, 0)
        x = 8.32 + i * .9
        max_h = .58
        h = max(.02, max_h * val / max_count)
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(3.77 - h), Inches(.42), Inches(h))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(YELLOW if impact in {"High", "Critical"} else "#7FC4F2")
        bar.line.fill.background()
        _ppt_text(slide, x - .07, 3.83, .58, .16, impact, 5.1, False, t["muted"])
        if val:
            _ppt_text(slide, x + .1, 3.58 - h, .25, .14, val, 5.5, True, t["text"])

    _ppt_text(slide, 8.05, 4.33, 3.5, .22, "Next confirmed deadlines", 11, True, t["text"])
    _ppt_box(slide, 8.05, 4.58, 4.93, .45, t["surface"], t["border"])
    deadlines = deadline_frame(df, horizon_days=365)
    if deadlines.empty:
        _ppt_text(slide, 8.22, 4.7, 4.5, .18, "No confirmed deadline in the next 12 months for this view.", 6.2, False, t["muted"])
    else:
        first = deadlines.iloc[0]
        _ppt_text(slide, 8.22, 4.67, 4.48, .25, f"{_text(first.get('compliance_deadline'))} · {_short(_scope_text(first), 24)} · {_short(first.get('english_title'), 38)}", 6.1, False, t["text"])

    # Divider and AI section.
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.35), Inches(5.25), Inches(12.63), Inches(.01))
    line.fill.solid(); line.fill.fore_color.rgb = _rgb(t["border"]); line.line.fill.background()
    _ppt_text(slide, .35, 5.44, 2.5, .22, "AI analysis", 11, True, t["text"])
    _ppt_text(slide, .35, 5.68, 10.8, .2, "Optional · grounded in visible structured records · informational only, not legal or tax advice", 6, False, t["muted"])
    ai_clean = _clean_ai(ai)
    if ai_clean:
        _ppt_box(slide, .35, 5.98, 12.63, .95, t["surface"], t["border"])
        _ppt_text(slide, .55, 6.12, 3.2, .16, "AI-GENERATED PORTFOLIO ANALYSIS", 6.5, True, PURPLE)
        _ppt_text(slide, .55, 6.34, 12.05, .45, ai_clean, 6.4, False, t["text"])
    else:
        _ppt_text(slide, .35, 6.1, 5, .2, "No AI analysis generated for this view yet.", 6.5, False, t["muted"])



def _ppt_finding_slide(slide, row, theme: str):
    t = _theme(theme)
    bg = slide.background.fill; bg.solid(); bg.fore_color.rgb = _rgb(t["bg"])
    _ppt_box(slide, .35, .25, 12.63, .75, t["surface"], t["border"])
    _ppt_text(slide, .55, .39, 4, .18, "VERIFIED CURRENT-WEEK FINDING", 7, True, PURPLE)
    _ppt_text(slide, .55, .59, 11.8, .30, _text(row.get("english_title")), 16, True, t["text"])
    fields = [
        ("Jurisdiction", row.get("jurisdiction")), ("Category", row.get("category")),
        ("Publication / update date", row.get("publication_update_date")), ("Effective / application date", row.get("effective_application_date")),
        ("Legislation", row.get("legislation")), ("Status", row.get("status")),
        ("Source", row.get("source_name")), ("Source URL", row.get("source_url")),
    ]
    for i, (label, value) in enumerate(fields):
        col = i % 2; r = i // 2; x = .45 + col*6.35; y = 1.18 + r*.55
        _ppt_text(slide, x, y, 2.8, .15, label.upper(), 5.8, True, t["muted"])
        _ppt_text(slide, x, y+.17, 5.9, .28, _text(value), 7, False, t["text"])
    _ppt_text(slide, .45, 3.47, 6, .22, "Scope — what products or services does it apply to?", 10, True, t["text"])
    _ppt_box(slide, .45, 3.76, 12.43, 1.0, t["surface"], t["border"])
    _ppt_text(slide, .65, 3.94, 12.0, .68, _text(row.get("scope")), 7.1, False, t["text"])
    _ppt_text(slide, .45, 4.98, 6, .22, "Summary — what changed and why does it matter?", 10, True, t["text"])
    _ppt_box(slide, .45, 5.28, 12.43, 1.02, t["surface"], t["border"])
    _ppt_text(slide, .65, 5.46, 12.0, .68, _text(row.get("summary")), 7.1, False, t["text"])
    _ppt_text(slide, .45, 6.53, 1.6, .18, "BUSINESS ACTION", 6.2, True, PURPLE)
    _ppt_text(slide, 1.95, 6.50, 10.7, .42, _text(row.get("business_action")), 7, False, t["text"])


def build_pptx(df: pd.DataFrame, ai_portfolio_analysis: str = "", theme: str = "light") -> bytes:
    """Export the current filtered weekly report plus one detail slide per finding."""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _ppt_dashboard_slide(slide, df, ai_portfolio_analysis, theme)
    for _, row in df.iterrows():
        detail = prs.slides.add_slide(prs.slide_layouts[6])
        _ppt_finding_slide(detail, row, theme)
    buf = BytesIO()
    prs.save(buf)
    return buf.getvalue()

